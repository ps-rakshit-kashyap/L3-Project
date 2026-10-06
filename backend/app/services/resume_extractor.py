import io
import re
from pathlib import Path
from sqlalchemy.orm import Session

import docx
import pypdf

from app.core.logging import logger
from app.models.resume import Resume
from app.services.storage import storage_service


class ResumeExtractionError(Exception):
    """Base exception for resume text extraction failures."""
    pass


class UnsupportedResumeFormatError(ResumeExtractionError):
    """Raised when the uploaded resume format is not supported."""
    pass


class CorruptResumeFileError(ResumeExtractionError):
    """Raised when a resume file is corrupted, encrypted, or unreadable."""
    pass


class EmptyResumeContentError(ResumeExtractionError):
    """Raised when the extracted text is empty or lacks meaningful content."""
    pass


class ResumeExtractorService:
    """
    Extracts and validates clean text from resumes in PDF, DOCX, and TXT formats.
    Caches extracted text on the Resume database entity to avoid duplicate computation.
    """

    MIN_TEXT_LENGTH = 20
    MIN_ALPHANUMERIC_CHARS = 10

    def extract_text_from_bytes(
        self,
        file_bytes: bytes,
        filename: str,
        content_type: str | None = None,
    ) -> str:
        """
        Extracts and validates text from raw bytes based on file extension and mime type.
        """
        if not file_bytes:
            raise EmptyResumeContentError("The provided resume file is 0 bytes.")

        ext = Path(filename).suffix.lower()
        mime = (content_type or "").lower()

        try:
            if ext == ".pdf" or "pdf" in mime:
                text = self._extract_pdf(file_bytes)
            elif ext in (".docx", ".doc") or "word" in mime or "officedocument" in mime:
                text = self._extract_docx(file_bytes)
            elif ext in (".txt", ".text") or "text/plain" in mime:
                text = self._extract_txt(file_bytes)
            else:
                raise UnsupportedResumeFormatError(
                    f"Unsupported file format '{ext}'. Supported formats: PDF, DOCX, TXT."
                )
        except (UnsupportedResumeFormatError, EmptyResumeContentError, CorruptResumeFileError):
            raise
        except Exception as exc:
            logger.error(f"Failed to parse resume '{filename}': {exc}")
            raise CorruptResumeFileError(f"Failed to parse resume document '{filename}': {exc}") from exc

        # Clean and validate meaningful content
        cleaned = self._clean_and_validate(text, filename)
        return cleaned

    def _extract_pdf(self, file_bytes: bytes) -> str:
        """Extracts text from PDF bytes using pypdf."""
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            if reader.is_encrypted:
                try:
                    # Attempt empty password decryption
                    reader.decrypt("")
                except Exception as exc:
                    raise CorruptResumeFileError("The PDF file is password protected.") from exc

            pages_text: list[str] = []
            for i, page in enumerate(reader.pages):
                try:
                    page_str = page.extract_text() or ""
                    if page_str.strip():
                        pages_text.append(page_str.strip())
                except Exception as exc:
                    logger.warning(f"Failed extracting text from PDF page {i}: {exc}")

            return "\n\n".join(pages_text)
        except Exception as exc:
            if isinstance(exc, (CorruptResumeFileError, EmptyResumeContentError)):
                raise
            raise CorruptResumeFileError(f"Corrupt or invalid PDF file: {exc}") from exc

    def _extract_docx(self, file_bytes: bytes) -> str:
        """Extracts text from DOCX bytes using python-docx."""
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            parts: list[str] = []

            # Paragraphs
            for p in doc.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    parts.append(p_text)

            # Tables (experience and education often inside tables)
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))

            return "\n".join(parts)
        except Exception as exc:
            raise CorruptResumeFileError(f"Corrupt or invalid DOCX document: {exc}") from exc

    def _extract_txt(self, file_bytes: bytes) -> str:
        """Decodes plain text bytes with fallback encodings."""
        for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                return file_bytes.decode(encoding)
            except UnicodeDecodeError:
                continue
        # Fallback with replacement
        return file_bytes.decode("utf-8", errors="replace")

    def _clean_and_validate(self, text: str, filename: str) -> str:
        """Normalizes extracted text and verifies sufficient meaningful content exists."""
        # Normalize carriage returns and excessive whitespace
        normalized = re.sub(r"\r\n?", "\n", text)
        normalized = re.sub(r"[ \t]+", " ", normalized)
        normalized = re.sub(r"\n{3,}", "\n\n", normalized).strip()

        if len(normalized) < self.MIN_TEXT_LENGTH:
            raise EmptyResumeContentError(
                f"Resume '{filename}' contains fewer than {self.MIN_TEXT_LENGTH} characters. "
                "The file may be blank, scanned image without text, or unreadable."
            )

        # Count alphanumeric characters
        alphanumeric_count = sum(1 for c in normalized if c.isalnum())
        if alphanumeric_count < self.MIN_ALPHANUMERIC_CHARS:
            raise EmptyResumeContentError(
                f"Resume '{filename}' does not contain sufficient readable text "
                f"({alphanumeric_count} alphanumeric characters found)."
            )

        return normalized

    def get_or_extract_text(self, db: Session, resume: Resume) -> str:
        """
        Retrieves extracted text for a resume entity.
        Returns cached text if present; otherwise downloads from storage, extracts,
        caches in the database, and returns the result.
        """
        if resume.extracted_text and len(resume.extracted_text.strip()) >= self.MIN_TEXT_LENGTH:
            return resume.extracted_text

        # Download raw bytes from storage
        file_bytes = storage_service.download_file(resume.file_path)

        # Extract & validate text
        text = self.extract_text_from_bytes(
            file_bytes=file_bytes,
            filename=resume.file_name,
            content_type=resume.file_type,
        )

        # Cache in database
        resume.extracted_text = text
        db.add(resume)
        db.commit()
        db.refresh(resume)

        return text


resume_extractor = ResumeExtractorService()
